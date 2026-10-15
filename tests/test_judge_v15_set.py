"""v1.5 LLM-as-judge — 기준 파일과 요약 날조 측정 세트·대응표·개발용 자료·채점 시트 생성 도구 테스트 (tools/judge_v15/)."""

import json
import sys
from collections import Counter
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools" / "judge_v15"))
import build_set as bs  # noqa: E402

from triage_eval.pipeline import judge  # noqa: E402
from triage_eval.pipeline.gate import GateError  # noqa: E402

DATA = ROOT / "data" / "judge_v15"


@pytest.fixture(scope="module")
def built():
    return bs.build()


@pytest.fixture(scope="module")
def committed():
    return {n: json.loads((DATA / n).read_text(encoding="utf-8")) for n in bs.FILES}


# ── 기준 파일 ────────────────────────────────────────────────

def test_criteria_match_plan():
    """계획 §6의 값과 같다."""
    c = judge.load_criteria()
    assert c["meta"]["judge_model"] == bs.JUDGE_MODEL == "gemma4:12b"
    assert c["natural"] == {"agreement_min": 90.0, "kappa_min": 0.61, "min_positive_for_kappa": 5, "false_alarm_max": 10.0}
    assert c["injected"] == {"recall_min": 90.0, "credit_requires_key": True}


def test_criteria_typo_is_rejected(tmp_path):
    bad = tmp_path / "c.toml"
    bad.write_text(judge.CRITERIA.read_text(encoding="utf-8").replace("recall_min", "recal_min"), encoding="utf-8")
    with pytest.raises(GateError, match="recal_min"):
        judge.load_criteria(bad)


# ── 측정 세트 (계획 §3) ──────────────────────────────────────

def test_committed_files_are_reproduced(built, committed):
    """같은 seed면 언제 만들어도 같은 세트다."""
    for name, payload in zip(bs.FILES, built):
        assert json.loads(bs.serialize(payload)) == committed[name]


def test_composition(built):
    """리포트 63개가 한 번씩(설계 의도 2): 자연 43 + 주입 20, 주입 종류는 기기4·수치4·시점3·추가 현상5·원인 단정4."""
    _, key, _ = built
    entries = list(key["map"].values())
    assert sorted(key["map"]) == [f"J{n:02d}" for n in range(1, 64)]
    assert sorted(e["item_id"] for e in entries) == sorted(bs.load_items())
    inj = [e for e in entries if e["injected"]]
    assert len(inj) == 20 and Counter(e["kind"] for e in inj) == {"기기": 4, "수치": 4, "시점": 3, "추가 현상": 5, "원인 단정": 4}
    assert all(e["log"] == bs.ADOPTED[0] for e in inj)


def test_no_summary_written_by_judge_model(built):
    """설계 의도 3 — 채점관(gemma4:12b)이 쓴 요약은 측정 세트·개발용 자료에 없다."""
    _, key, dev = built
    assert all(e["model"] != bs.JUDGE_MODEL for e in key["map"].values())
    assert all(d["model"] != bs.JUDGE_MODEL for d in dev["items"])
    assert {(e["log"], e["model"]) for e in key["map"].values()} <= set(bs.NATURAL_SOURCES)


def test_injected_facts_are_not_in_source(built):
    """설계 의도 5 — 주입한 사실의 핵심어는 주입 요약에만 있고, 시트에 보이는 원문 전체·원래 요약에는 없다."""
    measure, key, _ = built
    rows = {r["번호"]: r for r in measure["items"]}
    for jid, e in key["map"].items():
        if e["injected"]:
            shown = " ".join(str(v) for k, v in rows[jid].items() if k != "요약")
            assert all(k in e["summary"] and k not in shown and k not in e["original_summary"] for k in e["keys"])


def test_measure_set_hides_source_and_injection(built):
    """설계 의도 5 — 측정 세트에는 리포트 원문과 요약만 있다."""
    measure, key, _ = built
    allowed = {"번호", "입력 종류", "게시판 / 작성자", "게시 일시", "제목", "본문", "참고: 업데이트 공지", "참고: 기존 등록 이슈", "요약"}
    assert all(set(row) == allowed for row in measure["items"])
    text = json.dumps(measure, ensure_ascii=False)
    assert "injected" not in text and "item_id" not in text and "qwen" not in text and "gemma" not in text
    assert [r["요약"] for r in measure["items"]] == [key["map"][r["번호"]]["summary"] for r in measure["items"]]


def test_comma_does_not_reveal_injection(built):
    """설계 의도 4 — 쉼표 비율이 주입과 자연 응답에서 크게 다르지 않다 (처음 방식은 주입 20/20, 자연 1/80이었다).
    길이·숫자 단서는 남아 있고 계획 §8에 수치로 적었다."""
    _, key, _ = built
    inj = [e["summary"] for e in key["map"].values() if e["injected"]]
    nat = [e["summary"] for e in key["map"].values() if not e["injected"]]
    assert abs(sum("," in s for s in inj) / len(inj) - sum("," in s for s in nat) / len(nat)) <= 0.1


def test_stale_or_grounded_injection_is_refused():
    items = bs.load_items()
    base = bs.summaries(*bs.ADOPTED)
    i, (kind, original, injected, keys) = next(iter(bs.INJECTIONS.items()))
    with pytest.raises(ValueError, match="실행 기록과 다릅니다"):
        bs.check_injections({i: (kind, original + "!", injected, keys)}, items, base)
    word = items[i]["input"]["title"].split()[0]
    with pytest.raises(ValueError, match="날조가 아닙니다"):
        bs.check_injections({i: (kind, original, injected + " " + word, [word])}, items, base)


# ── 개발용 자료 (계획 §5, 설계 의도 6) ───────────────────────

def test_dev_set_is_separate(built):
    _, key, dev = built
    assert {d["log"] for d in dev["items"]} == {bs.DEV[0]}
    in_measure = {e["item_id"]: e["summary"] for e in key["map"].values()}
    assert all(d["summary"] != in_measure[d["item_id"]] for d in dev["items"])          # 같은 (리포트, 요약) 없음
    dev_inj = {d["item_id"] for d in dev["items"] if d["injected"]}
    assert len(dev_inj) == 10 and not dev_inj & {e["item_id"] for e in key["map"].values() if e["injected"]}
    dev_keys = {k for d in dev["items"] if d["injected"] for k in d["keys"]}
    assert not dev_keys & {k for e in key["map"].values() if e["injected"] for k in e["keys"]}


# ── 봉인 (설계 의도 7) ───────────────────────────────────────

def test_sealed_files_are_not_overwritten(monkeypatch, tmp_path, built):
    monkeypatch.setattr(bs, "OUT_DIR", tmp_path)
    payloads = dict(zip(bs.FILES, built))
    assert bs.write_sealed(payloads) == []                         # 없으면 쓴다
    assert bs.write_sealed(payloads) == []                         # 같으면 통과
    changed = dict(payloads)
    changed[bs.FILES[0]] = {**payloads[bs.FILES[0]], "seed": 0}
    before = (tmp_path / bs.FILES[0]).read_text(encoding="utf-8")
    assert bs.write_sealed(changed) == [bs.FILES[0]]               # 다르면 쓰지 않는다
    assert (tmp_path / bs.FILES[0]).read_text(encoding="utf-8") == before


# ── 채점 시트 ────────────────────────────────────────────────

def test_sheet_has_rows_and_dropdowns(built, tmp_path):
    measure, key, _ = built
    path = bs.write_sheet(measure, tmp_path)
    wb = load_workbook(path)
    ws = wb["판정"]
    assert [c.value for c in ws[1]] == bs.SHEET_COLS and ws.max_row == 64
    formulas = {dv.formula1 for dv in ws.data_validations.dataValidation}
    assert '"날조 있음,날조 없음"' in formulas and '"높음,보통,낮음"' in formulas
    cells = " ".join(str(c.value) for sheet in wb for row in sheet.iter_rows() for c in row if c.value)
    item_ids = {e["item_id"] for e in key["map"].values()}
    assert not any(f"{i} " in cells or cells.endswith(i) for i in item_ids if len(i) >= 3)   # 문항 ID 미노출
    assert "qwen" not in cells and "gemma" not in cells
