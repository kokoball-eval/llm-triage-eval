"""v1.5 LLM-as-judge — 채점관 실행·판정(pipeline/judge.py) 테스트. 가짜 채점관만 쓰고 실제 모델을 부르지 않는다.

측정 세트에 대한 실제 채점관 판정은 이 테스트에서 만들지 않는다(계획 §5 — 측정은 프롬프트·도구 커밋 뒤 1회).
"""

import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from triage_eval.pipeline import judge as jd
from triage_eval.pipeline.gate import GateError

ROOT = Path(__file__).resolve().parent.parent


class FakeJudge:
    """요약에 '가짜'가 있으면 날조 있음(문구 '가짜'), 아니면 없음. bad 번 만큼 먼저 깨진 출력을 낸다."""

    def __init__(self, bad: int = 0, load: float = 0.0, reload_at: str | None = None, raise_at: str | None = None):
        self.bad, self.load, self.reload_at, self.raise_at, self.prompts = bad, load, reload_at, raise_at, []

    def generate(self, model, prompt, options=None, think=None, format=None, keep_alive=None):
        if keep_alive == 0 or format is None:
            return {"response": "", "load_duration": 0}
        self.prompts.append(prompt)
        if self.raise_at and self.raise_at in prompt:
            raise ConnectionError("서버 연결 끊김")
        load = 16.4e9 if self.reload_at and self.reload_at in prompt else self.load
        if self.bad:
            self.bad -= 1
            return {"response": "날조 있음", "load_duration": load}
        summary = prompt.rsplit("[요약]\n", 1)[1]
        verdict = {"판정": jd.YES, "날조 문구": "가짜"} if "가짜" in summary else {"판정": jd.NO, "날조 문구": ""}
        return {"response": json.dumps(verdict, ensure_ascii=False), "load_duration": load}


def row(n: str, summary: str) -> dict:
    return {"번호": n, **{f: "" for f in jd.FIELDS}, "제목": "보스방 튕김", "본문": "보스방 들어가면 튕김", "요약": summary}


@pytest.fixture
def crit():
    return jd.load_criteria()


# ── 프롬프트·출력 읽기 ───────────────────────────────────────

def test_prompt_has_definition_and_all_fields():
    """설계 의도 2 — 계획 §2의 정의와 검토자 시트의 모든 원문 칸이 들어간다."""
    plan = (ROOT / "report/judge_v15_plan.md").read_text(encoding="utf-8")
    p = jd.build_prompt(row("J01", "보스방 튕김"))
    for line in ("원문을 다른 말로 바꿔 쓴 것, 원문 내용을 줄인 것은 날조가 아니다",
                 "추정임을 밝힌 표현은 날조가 아니다", "없거나 원문과 어긋나는 사실 진술"):
        assert line in p and line in plan
    assert all(f"{f}:" in p for f in jd.FIELDS) and p.rstrip().endswith("보스방 튕김")


@pytest.mark.parametrize("text,ok", [('{"판정": "날조 있음", "날조 문구": "x"}', True), ('{"판정": "아마도"}', False),
                                     ("날조 있음", False), ("", False)])
def test_parse_judgement(text, ok):
    assert (jd.parse_judgement(text) is not None) is ok


def test_unparsable_output_is_retried_once():
    """설계 의도 3 — 한 번 깨지면 다시 요청해 읽고, 두 번 깨지면 읽지 못함으로 둔다."""
    j = jd.judge_one(FakeJudge(bad=1), "fake", row("J01", "가짜 기기에서 튕김"))
    assert j["parsed"] and j["attempts"] == 2 and j["verdict"] == jd.YES
    j = jd.judge_one(FakeJudge(bad=2), "fake", row("J01", "가짜 기기에서 튕김"))
    assert not j["parsed"] and j["verdict"] is None


def test_mid_run_reload_is_detected():
    """설계 의도 4 — 워밍업 뒤 로드 시간이 RELOAD_SEC 를 넘은 응답을 찾아낸다(OBS-005)."""
    rows = [row("J01", "보스방 튕김"), row("J02", "다른 요약")]
    _, reloads = jd.run_judge(FakeJudge(reload_at="다른 요약"), "fake", rows)
    assert reloads and reloads[0].startswith("J02")


# ── 지표·판정 (계획 §6) ──────────────────────────────────────

def make(n_nat_agree: int, n_nat_disagree_fa: int, inj_hit: int, inj_wrong_phrase: int = 0, human_yes_nat: int = 0):
    """자연 43(일치 n_nat_agree, 거짓 경보 n_nat_disagree_fa) + 주입 20(문구 맞게 검출 inj_hit, 다른 문구 inj_wrong_phrase)."""
    key, human, judged = {}, {}, {}
    n = 0
    for k in range(43):
        n += 1
        j = f"J{n:02d}"
        key[j] = {"injected": False, "summary": "s", "item_id": f"N{k:02d}"}
        hy = k < human_yes_nat
        human[j] = {"verdict": jd.YES if hy else jd.NO, "phrase": ""}
        if k < n_nat_agree:
            judged[j] = {"verdict": human[j]["verdict"], "phrase": "", "parsed": True}
        else:
            judged[j] = {"verdict": jd.YES if not hy else jd.NO, "phrase": "", "parsed": True}
    for k in range(20):
        n += 1
        j = f"J{n:02d}"
        key[j] = {"injected": True, "keys": ["가짜"], "kind": "기기", "summary": "s 가짜", "item_id": f"I{k:02d}"}
        human[j] = {"verdict": jd.YES, "phrase": "가짜"}
        if k < inj_hit:
            judged[j] = {"verdict": jd.YES, "phrase": "가짜", "parsed": True}
        elif k < inj_hit + inj_wrong_phrase:
            judged[j] = {"verdict": jd.YES, "phrase": "다른 곳", "parsed": True}
        else:
            judged[j] = {"verdict": jd.NO, "phrase": "", "parsed": True}
    return key, human, judged


def test_boundary_values_pass(crit):
    """기준값과 같으면 충족 — 일치 39/43(90.7%), 주입 18/20(90%), 거짓 경보 4/43(9.3%)."""
    key, human, judged = make(39, 4, 18)
    mx = jd.metrics(key, human, judged, crit)
    assert mx["injected_recall"] == 90.0 and not mx["kappa_used"]
    assert jd.decide(mx, crit)["adopted"]


def test_wrong_phrase_is_not_counted_as_detection(crit):
    """설계 의도 5 — 날조 있음이라도 다른 곳을 지목하면 검출이 아니다(문구 무관 검출률에는 센다)."""
    key, human, judged = make(43, 0, 17, inj_wrong_phrase=3)
    mx = jd.metrics(key, human, judged, crit)
    assert mx["injected_detected"] == 17 and mx["injected_detected_any"] == 20
    assert not jd.decide(mx, crit)["adopted"]


def test_false_alarms_fail(crit):
    key, human, judged = make(38, 5, 20)
    mx = jd.metrics(key, human, judged, crit)
    assert mx["false_alarm_rate"] > 10.0 and not jd.decide(mx, crit)["adopted"]


def test_kappa_is_used_with_enough_positives(crit):
    """검토자 있음이 5건 이상이면 κ도 판정에 쓴다 — 있음 5건을 모두 놓치면 일치율은 높아도 κ로 미채택."""
    key, human, judged = make(43, 0, 20, human_yes_nat=5)
    for j in ("J01", "J02", "J03", "J04", "J05"):        # 검토자 있음 5건을 채점관이 모두 놓침
        judged[j] = {"verdict": jd.NO, "phrase": "", "parsed": True}
    mx = jd.metrics(key, human, judged, crit)
    assert mx["natural_agreement"] == pytest.approx(88.37, abs=0.01)
    assert mx["kappa_used"] and mx["natural_judge_detected"] == 0
    names = {c["name"]: c["passed"] for c in jd.decide(mx, crit)["checks"]}
    assert names["자연 응답 κ"] is False


def test_unparsed_counts_as_disagreement_and_miss(crit):
    """계획 §4 — 두 번 읽지 못한 건은 불일치·놓침."""
    key, human, judged = make(43, 0, 20)
    judged["J01"] = {"verdict": None, "phrase": "", "parsed": False}
    judged["J63"] = {"verdict": None, "phrase": "", "parsed": False}
    mx = jd.metrics(key, human, judged, crit)
    assert mx["natural_agree"] == 42 and mx["injected_detected"] == 19 and mx["unparsed"] == 2


# ── 측정 모드 안전장치 (설계 의도 6) ─────────────────────────

def _blank_sheet(tmp_path):
    import sys
    sys.path.insert(0, str(ROOT / "tools" / "judge_v15"))
    import build_set as bs
    measure = json.loads((ROOT / "data/judge_v15/measure_set_v15.json").read_text(encoding="utf-8"))
    return bs.write_sheet(measure, tmp_path)


def test_incomplete_sheet_stops_before_judge(tmp_path, crit):
    """시트에 빈칸이 있으면 채점관을 한 번도 부르지 않고 멈춘다."""
    path = _blank_sheet(tmp_path)
    fake = FakeJudge()
    with pytest.raises(GateError, match="빈칸"):
        jd.run_measure(fake, crit, path, guard=False)
    assert fake.prompts == []


def test_complete_sheet_is_read(tmp_path):
    path = _blank_sheet(tmp_path)
    wb = load_workbook(path)
    ws = wb["판정"]
    col = [c.value for c in ws[1]].index("판정") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=col, value=jd.NO)
    wb.save(path)
    key = json.loads((ROOT / "data/judge_v15/answer_key_v15.json").read_text(encoding="utf-8"))["map"]
    human = jd.read_rater_sheet(path, list(key))
    assert len(human) == 63 and all(h["verdict"] == jd.NO for h in human.values())


def test_dev_mode_never_reads_measure_set(monkeypatch, tmp_path, crit):
    """개발용 모드는 개발용 자료만 채점관에 넣는다 (계획 §5)."""
    monkeypatch.setattr(jd, "DEV_RESULT", tmp_path / "dev.json")
    fake = FakeJudge()
    monkeypatch.setattr(jd, "get_model_digests", lambda c, ms: {m: "fake" for m in ms})
    assert jd.run_dev(fake, crit) == 0
    measure = json.loads((ROOT / "data/judge_v15/measure_set_v15.json").read_text(encoding="utf-8"))
    dev = json.loads((ROOT / "data/judge_v15/dev_set_v15.json").read_text(encoding="utf-8"))
    sent = "\n".join(fake.prompts)
    assert len(fake.prompts) == len(dev["items"])
    only_measure = {r["요약"] for r in measure["items"]} - {d["summary"] for d in dev["items"]}
    assert not any(f"[요약]\n{s}\n" in sent for s in only_measure)


def test_committed_dev_result_matches_prompt():
    """커밋된 개발용 결과가 지금 프롬프트로 만든 것이다."""
    dev = json.loads((ROOT / "data/judge_v15/dev_result_v15.json").read_text(encoding="utf-8"))
    assert dev["prompt_version"] == jd.PROMPT_VERSION and dev["prompt_sha256"] == jd.prompt_fingerprint()
    assert dev["options"] == jd.OPTIONS and dev["reloads"] == []


# ── 독립 검토 반영 ───────────────────────────────────────────

def test_unparsed_on_rater_no_counts_as_false_alarm(crit):
    """설계 의도 10 — 검토자가 없음으로 본 자연 응답을 읽지 못하면 거짓 경보로도 센다(가장 불리하게)."""
    key, human, judged = make(43, 0, 20)
    judged["J01"] = {"verdict": None, "phrase": "", "parsed": False}
    mx = jd.metrics(key, human, judged, crit)
    assert mx["false_alarm"] == 1


def test_zero_denominator_fails_false_alarm_check(crit):
    """검토자가 자연 응답을 모두 있음으로 보면 거짓 경보율 값이 없고, 그 기준은 충족하지 못한 것으로 본다."""
    key, human, judged = make(43, 0, 20, human_yes_nat=43)
    mx = jd.metrics(key, human, judged, crit)
    names = {c["name"]: c["passed"] for c in jd.decide(mx, crit)["checks"]}
    assert mx["false_alarm_rate"] is None and names["거짓 경보율(%)"] is False


def test_near_boundary_flag(crit):
    key, human, judged = make(39, 4, 20)
    assert jd.metrics(key, human, judged, crit)["near_boundary"]
    key, human, judged = make(43, 0, 20)
    assert not jd.metrics(key, human, judged, crit)["near_boundary"]


def test_require_key_false_counts_any_yes():
    j = {"verdict": jd.YES, "phrase": "다른 곳"}
    assert jd.detected(j, ["가짜"], False) and not jd.detected(j, ["가짜"], True)


def _filled_sheet(tmp_path, value=jd.NO):
    path = _blank_sheet(tmp_path)
    wb = load_workbook(path)
    ws = wb["판정"]
    col = [c.value for c in ws[1]].index("판정") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=col, value=value)
    wb.save(path)
    return path


def test_measure_refuses_no_write(tmp_path, crit):
    with pytest.raises(GateError, match="no-write"):
        jd.run_measure(FakeJudge(), crit, _filled_sheet(tmp_path), write=False)


def test_measure_full_flow_with_fake_judge(monkeypatch, tmp_path, crit):
    """채운 시트 + 가짜 채점관으로 지표 → 판정 → 보고서까지 끝난다 (측정은 한 번뿐이라 마지막 단계가 죽으면 안 된다)."""
    monkeypatch.setattr(jd, "RESULT_JSON", tmp_path / "result.json")
    monkeypatch.setattr(jd, "REPORT", tmp_path / "report.md")
    monkeypatch.setattr(jd, "get_model_digests", lambda c, ms: {m: "fake" for m in ms})
    code = jd.run_measure(FakeJudge(), crit, _filled_sheet(tmp_path), guard=False)
    assert code in (0, 1)
    result = json.loads((tmp_path / "result.json").read_text(encoding="utf-8"))
    assert result["metrics"]["natural_n"] == 42 and result["metrics"]["injected_n"] == 20   # A06 제외 (ISSUE-012)
    assert [x["item_id"] for x in result["metrics"]["excluded"]] == ["A06"]
    assert "## 판정:" in (tmp_path / "report.md").read_text(encoding="utf-8")


def test_measure_saves_result_before_printing(monkeypatch, tmp_path, crit):
    """설계 의도 12 (ISSUE-014) — 화면 출력이 실패해도 측정 결과는 이미 파일로 남아 있다."""
    monkeypatch.setattr(jd, "RESULT_JSON", tmp_path / "result.json")
    monkeypatch.setattr(jd, "REPORT", tmp_path / "report.md")
    monkeypatch.setattr(jd, "get_model_digests", lambda c, ms: {m: "fake" for m in ms})
    sheet = _filled_sheet(tmp_path)

    def failing_print(*args, **kwargs):
        if "✅" in str(args) or "❌" in str(args):
            raise UnicodeEncodeError("cp949", "✅", 0, 1, "illegal multibyte sequence")

    monkeypatch.setattr("builtins.print", failing_print)
    with pytest.raises(UnicodeEncodeError):
        jd.run_measure(FakeJudge(), crit, sheet, guard=False)
    assert json.loads((tmp_path / "result.json").read_text(encoding="utf-8"))["metrics"]["natural_n"] == 42
    assert (tmp_path / "report.md").exists()


def test_cli_prints_on_cp949_console(monkeypatch):
    """설계 의도 12 (ISSUE-014) — cp949 콘솔에서도 ✅·한글 출력이 멈추지 않는다."""
    import io
    out = io.TextIOWrapper(io.BytesIO(), encoding="cp949")
    monkeypatch.setattr(sys, "stdout", out)
    monkeypatch.setattr(sys, "stderr", io.TextIOWrapper(io.BytesIO(), encoding="cp949"))   # pytest의 stderr를 바꾸지 않게
    assert jd.main(["--measure"]) == 2
    print("✅")
    out.flush()
    assert out.encoding == "utf-8" and "✅" in out.buffer.getvalue().decode("utf-8")


def test_measure_with_connection_error_is_invalid(monkeypatch, tmp_path, crit):
    """설계 의도 9 — 통신 오류가 나면 판정하지 않고 무효 실행을 기록한다."""
    monkeypatch.setattr(jd, "RESULT_JSON", tmp_path / "result.json")
    monkeypatch.setattr(jd, "INVALID_RUNS", tmp_path / "invalid.json")
    monkeypatch.setattr(jd, "get_model_digests", lambda c, ms: {m: "fake" for m in ms})
    measure = json.loads((ROOT / "data/judge_v15/measure_set_v15.json").read_text(encoding="utf-8"))
    with pytest.raises(GateError, match="측정 환경 오염"):
        jd.run_measure(FakeJudge(raise_at=measure["items"][5]["요약"]), crit, _filled_sheet(tmp_path), guard=False)
    assert not (tmp_path / "result.json").exists()
    assert "통신 오류" in json.loads((tmp_path / "invalid.json").read_text(encoding="utf-8"))[0]["reasons"][0]


def test_measure_guard_refuses_existing_result_and_uncommitted_sheet(monkeypatch, tmp_path, crit):
    """설계 의도 8 — 결과 파일이 이미 있거나 시트가 커밋되지 않았으면 측정하지 않는다."""
    monkeypatch.setattr(jd, "get_model_digests", lambda c, ms: {m: "fake" for m in ms})
    existing = tmp_path / "result.json"
    existing.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(jd, "RESULT_JSON", existing)
    with pytest.raises(GateError) as e:
        jd.measure_guard(FakeJudge(), crit, _filled_sheet(tmp_path))
    assert "이미 있습니다" in str(e.value) and "커밋되어 있지 않습니다" in str(e.value) and "digest" in str(e.value)


@pytest.mark.parametrize("mutate,expected", [
    (lambda ws, col: ws.cell(row=3, column=1, value=ws.cell(row=2, column=1).value), "중복"),
    (lambda ws, col: ws.cell(row=2, column=[c.value for c in ws[1]].index("요약") + 1, value="바뀐 요약"), "요약이 측정 세트와 다릅니다"),
])
def test_sheet_anomalies_are_refused(tmp_path, mutate, expected):
    path = _filled_sheet(tmp_path)
    wb = load_workbook(path)
    ws = wb["판정"]
    mutate(ws, None)
    wb.save(path)
    measure = json.loads((ROOT / "data/judge_v15/measure_set_v15.json").read_text(encoding="utf-8"))
    with pytest.raises(GateError, match=expected):
        jd.read_rater_sheet(path, [r["번호"] for r in measure["items"]], {r["번호"]: r["요약"] for r in measure["items"]})


def test_cli_measure_without_sheet_exits_2():
    assert jd.main(["--measure"]) == 2


def test_excluded_item_is_removed_from_natural_metrics(crit):
    """설계 의도 11 — 제외 문항은 자연 응답 지표에서 빠지고 판정은 따로 남는다 (ISSUE-012)."""
    key, human, judged = make(43, 0, 20)
    key["J01"]["item_id"] = "A06"
    judged["J01"] = {"verdict": jd.YES, "phrase": "x", "parsed": True}     # 불일치지만 지표에 들어가지 않는다
    mx = jd.metrics(key, human, judged, crit)
    assert mx["natural_n"] == 42 and mx["natural_agreement"] == 100.0 and mx["false_alarm"] == 0
    assert mx["excluded"] == [{"번호": "J01", "item_id": "A06", "human": jd.NO, "judge": jd.YES}]


def test_excluded_item_must_be_natural(crit):
    key, human, judged = make(43, 0, 20)
    key["J63"]["item_id"] = "A06"
    with pytest.raises(GateError, match="주입 문항"):
        jd.metrics(key, human, judged, crit)


def test_committed_excluded_item_is_natural_in_measure_set(crit):
    """기준 파일의 제외 문항이 실제 측정 세트에서 자연 응답이고, 빼면 자연 응답이 42건이다."""
    key = json.loads((ROOT / "data/judge_v15/answer_key_v15.json").read_text(encoding="utf-8"))["map"]
    ex = [e for e in key.values() if e["item_id"] in crit["natural"]["excluded_items"]]
    assert len(ex) == 1 and not ex[0]["injected"]
    assert sum(not e["injected"] and e["item_id"] not in crit["natural"]["excluded_items"] for e in key.values()) == 42
