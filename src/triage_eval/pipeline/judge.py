"""v1.5 LLM-as-judge — 요약 날조 판정의 채점관(gemma4:12b)과 검토자 채점의 일치를 잰다.

계획(사전 등록): report/judge_v15_plan.md · 기준: judge_v15.toml · 세트: tools/judge_v15/build_set.py

사용법
------
    uv run triage-judge --dev                                   # 개발용 자료로 채점관 실행 (프롬프트 개발, 계획 §5)
    uv run triage-judge --measure --sheet <채점한 시트.xlsx>     # 측정 1회 (계획 §5-4)

종료 코드
---------
    --dev       0  실행 완료   2  측정 환경 오염(실행 도중 모델 재로드)
    --measure   0  채택   1  미채택   2  판정 불가 (시트 빈칸·번호 불일치, 측정 환경 오염, 기준 파일 오류)

출력
----
    data/judge_v15/dev_result_v15.json     개발용 실행 결과 (마지막 실행만 남는다)
    data/judge_v15/judge_result_v15.json   측정 결과 (검토자 판정·채점관 판정·지표·판정)
    report/judge_v15.md                    사람이 읽는 판정표

[설계 의도]
1. 기준 파일에 모르는 섹션·키가 있거나 필요한 키가 없으면 읽지 않는다(v1.1.1 ISSUE-003). 오타로 기준이 조용히 꺼진 채
   채택 판정이 나는 것을 막는다.
2. 채점관에게는 검토자 시트와 같은 칸(입력 종류·게시판/작성자·게시 일시·제목·본문·공지·기존 이슈·요약)과 계획 §2의 날조 정의를
   같은 문구로 준다. 사람과 채점관이 다른 정보로 판정하면 일치율이 정의가 아니라 정보 차이를 잰다.
3. 출력은 JSON 스키마(판정 enum + 날조 문구)로 제한한다(Ollama 구조화 출력). 자유 서술을 파싱하면 형식 오류가 판정 오류로 섞인다.
   그래도 읽지 못하면 1번 다시 요청하고, 또 실패하면 검토자와 불일치·주입 놓침으로 센다(계획 §4, 가장 불리하게).
4. temperature 0, seed 1, think=False. 같은 입력이면 같은 판정이 나와야 측정을 다시 확인할 수 있다. 실행 전에 채점관을 메모리에서
   내리고 워밍업 1회 뒤, 워밍업을 뺀 응답의 로드 시간이 RELOAD_SEC 를 넘으면 실행 도중 재로드로 보고 판정하지 않는다(OBS-005).
5. 주입 날조는 채점관이 날조 있음으로 판정하고 지목한 문구에 주입 핵심어가 들어 있어야 검출로 센다(judge_v15.toml
   credit_requires_key). 원래 요약의 다른 부분을 지목해 우연히 맞힌 것을 검출로 세지 않기 위해서다. 문구와 무관한 검출률도 함께 기록한다.
6. 측정 모드는 개발용 모드와 명령을 나눈다(--measure). 측정 세트를 프롬프트 개발 중에 실수로 돌리지 않게 하기 위해서다.
   검토자 시트가 63행 모두 채워지지 않았으면 채점관을 실행하기 전에 멈춘다.
7. 프롬프트는 개발용 자료(seed 11)로만 다듬어 j1로 확정했다(계획 §5). 개발 경과와 개발용 판정 내용은 검토자 채점이 끝난 뒤
   결과 보고서에 적는다. 검토자가 채점 중에 채점관의 판정 경향을 알면 사람 판정이 그쪽으로 기운다.
8. 측정은 한 번만 한다(계획 §5-4). 측정 모드는 결과 파일이 이미 있으면, --no-write 이면, 프롬프트·기준·측정 세트 파일에
   커밋하지 않은 변경이 있거나 검토자 시트가 커밋되지 않았으면, 프롬프트 지문·모델 digest가 개발용 실행과 다르면 실행하지 않는다.
9. 통신 오류(예외)는 "출력을 읽지 못함"과 구분한다. 측정 도중 통신 오류나 모델 재로드가 나면 측정 환경 오염으로 보고 판정하지 않으며,
   무효 실행을 data/judge_v15/invalid_runs_v15.json 에 남긴다(계획 §4 — 다시 잰 사실과 이유를 기록).
10. 읽지 못한 건은 가장 불리하게 센다: 일치율·κ에서는 불일치(κ에서는 세 번째 범주), 주입은 놓침, 검토자가 없음으로 본 자연 응답이면
    거짓 경보. 검토자가 자연 응답을 모두 있음으로 보면 거짓 경보율의 분모가 0이라 값이 없고, 그 기준은 충족하지 못한 것으로 본다.
11. 기준 파일의 excluded_items 문항은 자연 응답 지표(일치율·κ·거짓 경보)에서 뺀다. 검토자 채점 전에 그 문항의 날조 내용이
    검토자에게 노출되어 사람 판정이 독립적이지 않기 때문이다(docs/issue_log.md ISSUE-012). 뺀 문항의 두 판정은 결과에 따로 남긴다.
"""

import argparse
import json
import subprocess
import sys
import time
import tomllib
from collections import Counter
from pathlib import Path

from triage_eval.common.detect_hallucination import detect
from triage_eval.common.ollama_runtime import get_model_digests, sha256_text
from triage_eval.common.paths import ROOT
from triage_eval.pipeline.agreement import cohen_kappa
from triage_eval.pipeline.gate import GateError
from triage_eval.pipeline.score import rel

CRITERIA = ROOT / "judge_v15.toml"
KNOWN = {
    "meta": {"version", "judge_model", "rater", "measure_set", "answer_key", "dev_set"},
    "natural": {"agreement_min", "kappa_min", "min_positive_for_kappa", "false_alarm_max", "excluded_items"},
    "injected": {"recall_min", "credit_requires_key"},
}
DEV_RESULT = ROOT / "data" / "judge_v15" / "dev_result_v15.json"
RESULT_JSON = ROOT / "data" / "judge_v15" / "judge_result_v15.json"
INVALID_RUNS = ROOT / "data" / "judge_v15" / "invalid_runs_v15.json"   # 설계 의도 9
GUARDED = ["src/triage_eval/pipeline/judge.py", "judge_v15.toml", "data/judge_v15/measure_set_v15.json",
           "data/judge_v15/answer_key_v15.json", "data/judge_v15/dev_result_v15.json"]   # 설계 의도 8
REPORT = ROOT / "report" / "judge_v15.md"

YES, NO = "날조 있음", "날조 없음"
OPTIONS = {"temperature": 0, "seed": 1, "num_predict": 200, "num_ctx": 4096}   # 설계 의도 4
THINK = False
RELOAD_SEC = 1.0
SCHEMA = {"type": "object", "properties": {"판정": {"type": "string", "enum": [YES, NO]}, "날조 문구": {"type": "string"}},
          "required": ["판정", "날조 문구"]}   # 설계 의도 3
FIELDS = ["입력 종류", "게시판 / 작성자", "게시 일시", "제목", "본문", "참고: 업데이트 공지", "참고: 기존 등록 이슈"]

PROMPT_VERSION = "j1"   # 설계 의도 7
NONE_TEXT = "(없음)"
PROMPT = """너는 게임 버그 리포트 트리아지 결과를 검수하는 QA 검토자다. 아래 [리포트 원문]과 모델이 쓴 [요약]을 비교해, 요약에 날조가 있는지 판정한다.

[날조의 정의]
- 요약에 리포트 원문(게시판·게시 일시·제목·본문·업데이트 공지·기존 등록 이슈)에 없거나 원문과 어긋나는 사실 진술이 하나라도 있으면 "날조 있음"이다.
- 원문을 다른 말로 바꿔 쓴 것, 원문 내용을 줄인 것은 날조가 아니다.
- "~로 추정", "~일 수 있음"처럼 추정임을 밝힌 표현은 날조가 아니다.

[판정 방법]
요약을 사실 진술 단위로 나누고, 각 진술이 리포트 원문에 근거가 있는지 하나씩 확인한다.

[출력]
JSON으로만 답한다. "판정"은 "날조 있음" 또는 "날조 없음". "날조 문구"는 날조 있음이면 요약에서 근거 없는 부분을 그대로 옮겨 적고, 날조 없음이면 빈 문자열로 둔다.
"""


# ── 기준 파일 ────────────────────────────────────────────────

def load_criteria(path: Path = CRITERIA) -> dict:
    """설계 의도 1."""
    if not path.exists():
        raise GateError(f"기준 파일이 없습니다: {path}")
    c = tomllib.loads(path.read_text(encoding="utf-8"))
    problems = [f"모르는 섹션 [{s}]" for s in c if s not in KNOWN]
    problems += [f"[{s}] 모르는 키 {k!r}" for s, body in c.items() if s in KNOWN for k in body if k not in KNOWN[s]]
    problems += [f"[{s}] {k} 가 없습니다" for s, keys in KNOWN.items() for k in keys if k not in c.get(s, {})]
    if problems:
        raise GateError("기준 파일 오류 → " + "; ".join(problems))
    return c


# ── 채점관 ───────────────────────────────────────────────────

def source_of(row: dict) -> str:
    return "\n".join(f"{f}: {row.get(f) or NONE_TEXT}" for f in FIELDS)


def build_prompt(row: dict) -> str:
    """설계 의도 2."""
    return f"{PROMPT}\n[리포트 원문]\n{source_of(row)}\n\n[요약]\n{row['요약']}\n"


def parse_judgement(text: str) -> dict | None:
    try:
        data = json.loads(text)
    except (json.JSONDecodeError, TypeError):
        return None
    if not isinstance(data, dict) or data.get("판정") not in (YES, NO):
        return None
    return {"verdict": data["판정"], "phrase": str(data.get("날조 문구") or "").strip()}


def judge_one(client, model: str, row: dict) -> dict:
    """설계 의도 3 — 1회 + 읽지 못하면 1회 재요청."""
    loads, raws, errors = [], [], 0
    for attempt in (1, 2):
        start = time.perf_counter()
        try:
            r = client.generate(model=model, prompt=build_prompt(row), options=OPTIONS, think=THINK, format=SCHEMA)
            raw, load = r.get("response", ""), round(r.get("load_duration", 0) / 1e9, 3)
        except Exception as e:   # 설계 의도 9 — 통신 오류는 따로 센다
            raw, load, errors = f"ERROR: {e}", 0.0, errors + 1
        raws.append(raw)
        loads.append(load)
        parsed = parse_judgement(raw)
        if parsed:
            return {**parsed, "parsed": True, "attempts": attempt, "raw": raws, "load_duration_sec": loads,
                    "errors": errors, "elapsed_sec": round(time.perf_counter() - start, 3)}
    return {"verdict": None, "phrase": "", "parsed": False, "attempts": 2, "raw": raws, "load_duration_sec": loads,
            "errors": errors}


def run_judge(client, model: str, rows: list[dict]) -> tuple[list[dict], list[str]]:
    """설계 의도 4 — (판정 목록, 재로드 감지 목록)."""
    try:
        client.generate(model=model, prompt="", keep_alive=0)
    except Exception:
        pass
    client.generate(model=model, prompt="워밍업", options=OPTIONS, think=THINK)
    out, reloads = [], []
    for row in rows:
        j = judge_one(client, model, row)
        if any(x > RELOAD_SEC for x in j["load_duration_sec"]):
            reloads.append(f"{row['번호']} (로드 {max(j['load_duration_sec'])}초)")
        if j["errors"]:
            reloads.append(f"{row['번호']} (통신 오류 {j['errors']}회)")
        out.append({"번호": row["번호"], **j})
        print(f"  {row['번호']}: {j['verdict'] or '읽지 못함'} {j['phrase']}")
    return out, reloads


def detected(j: dict, keys: list[str], require_key: bool) -> bool:
    """설계 의도 5."""
    if j["verdict"] != YES:
        return False
    return any(k in j["phrase"] for k in keys) if require_key else True


def prompt_fingerprint() -> str:
    """프롬프트 문구·입력 칸·빈칸 표시·출력 스키마를 함께 덮는 지문."""
    return sha256_text(json.dumps([PROMPT, FIELDS, NONE_TEXT, SCHEMA], ensure_ascii=False, sort_keys=True))


def run_meta(client, model: str, digest_before: str | None) -> dict:
    return {"judge_model": model, "model_digest": digest_before,
            "model_digest_after": get_model_digests(client, [model]).get(model),
            "prompt_version": PROMPT_VERSION, "prompt_sha256": prompt_fingerprint(), "options": OPTIONS, "think": THINK,
            "executed_at": time.strftime("%Y-%m-%dT%H:%M:%S")}


# ── 개발용 모드 (계획 §5) ────────────────────────────────────

def dev_rows(dev: dict) -> list[dict]:
    return [{"번호": f"D{n:02d}", **{f: d.get(f, "") for f in FIELDS}, "요약": d["summary"], "_item": d}
            for n, d in enumerate(dev["items"], start=1)]


def dev_summary(rows: list[dict], results: list[dict]) -> dict:
    """개발용 자료에는 사람 판정이 없으므로, 주입 검출과 자연 응답의 있음 판정 목록만 정리한다(사람이 직접 읽고 판단)."""
    inj = [(r, j) for r, j in zip(rows, results) if r["_item"]["injected"]]
    nat = [(r, j) for r, j in zip(rows, results) if not r["_item"]["injected"]]
    return {
        "n": len(rows), "unparsed": sum(not j["parsed"] for j in results),
        "injected_detected_key": sum(detected(j, r["_item"]["keys"], True) for r, j in inj),
        "injected_detected_any": sum(j["verdict"] == YES for _, j in inj), "injected_n": len(inj),
        "injected_missed": [{"item_id": r["_item"]["item_id"], "kind": r["_item"]["kind"], "summary": r["요약"],
                             "verdict": j["verdict"], "phrase": j["phrase"]}
                            for r, j in inj if not detected(j, r["_item"]["keys"], True)],
        "natural_yes": [{"item_id": r["_item"]["item_id"], "summary": r["요약"], "phrase": j["phrase"]}
                        for r, j in nat if j["verdict"] == YES],
        "natural_n": len(nat),
    }


def run_dev(client, crit: dict, write: bool = True) -> int:
    dev = json.loads((ROOT / crit["meta"]["dev_set"]).read_text(encoding="utf-8"))
    rows = dev_rows(dev)
    model = crit["meta"]["judge_model"]
    digest = get_model_digests(client, [model]).get(model)
    results, reloads = run_judge(client, model, rows)
    summary = dev_summary(rows, results)
    print(f"\n개발용 {summary['n']}건 · 읽지 못함 {summary['unparsed']} · 주입 검출(문구 일치) "
          f"{summary['injected_detected_key']}/{summary['injected_n']} · 자연 응답 있음 판정 {len(summary['natural_yes'])}/{summary['natural_n']}")
    if reloads:
        print("측정 환경 오염 — 실행 도중 모델 재로드: " + ", ".join(reloads))
    if write:
        DEV_RESULT.write_text(json.dumps({**run_meta(client, model, digest), "reloads": reloads, "summary": summary,
                                          "results": [{**j, "item_id": r["_item"]["item_id"]} for r, j in zip(rows, results)]},
                                         ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"저장: {rel(DEV_RESULT)}")
    return 2 if reloads else 0


# ── 측정 모드 (계획 §5-4, §6) ────────────────────────────────

def read_rater_sheet(path: Path, expected: list[str], summaries: dict[str, str] | None = None) -> dict[str, dict]:
    """설계 의도 6 — {번호: {verdict, phrase, 확신도, 메모}}. 빈칸·선택지 밖 값·번호 불일치면 GateError."""
    from openpyxl import load_workbook   # 개발 의존성

    if not path.exists():
        raise GateError(f"채점 시트가 없습니다: {path}")
    wb = load_workbook(path, data_only=True)
    if "판정" not in wb.sheetnames:
        raise GateError(f"{path.name} 에 '판정' 시트가 없습니다")
    ws = wb["판정"]
    header = [c.value for c in ws[1]]
    out, problems = {}, []
    for row in ws.iter_rows(min_row=2, values_only=True):
        v = dict(zip(header, row))
        jid = str(v.get("번호") or "").strip()
        if not jid:
            continue
        if jid in out:
            problems.append(f"{jid} 번호 중복")
        verdict = str(v.get("판정") or "").strip()
        if verdict not in (YES, NO):
            problems.append(f"{jid} 판정 {'빈칸' if not verdict else repr(verdict)}")
        if summaries is not None and str(v.get("요약") or "").strip() != summaries.get(jid, "").strip():
            problems.append(f"{jid} 요약이 측정 세트와 다릅니다")
        out[jid] = {"verdict": verdict, "phrase": str(v.get("날조 문구") or "").strip(),
                    "확신도": v.get("확신도"), "메모": v.get("메모")}
    if sorted(out) != sorted(expected):
        problems.append("시트의 번호가 측정 세트와 다릅니다")
    if problems:
        raise GateError(f"채점 시트에 고칠 곳이 {len(problems)}곳 있습니다 → " + "; ".join(problems))
    return out


def pct(a: int, b: int) -> float | None:
    return round(a / b * 100, 2) if b else None


def metrics(key: dict, human: dict, judged: dict, crit: dict) -> dict:
    """계획 §6."""
    excluded = set(crit["natural"]["excluded_items"])                                              # 설계 의도 11
    nat = [j for j, e in key.items() if not e["injected"] and e["item_id"] not in excluded]
    inj = [j for j, e in key.items() if e["injected"]]
    if any(key[j]["item_id"] in excluded for j in inj):
        raise GateError("제외 문항이 주입 문항입니다 — 자연 응답만 뺄 수 있습니다")
    h = {j: human[j]["verdict"] for j in key}
    m = {j: judged[j]["verdict"] or ("불일치" if not judged[j]["parsed"] else None) for j in key}   # 읽지 못함 → 불일치
    agree = sum(h[j] == m[j] for j in nat)
    human_yes = [j for j in nat if h[j] == YES]
    human_no = [j for j in nat if h[j] == NO]
    kappa = cohen_kappa([h[j] for j in nat], [m[j] if m[j] in (YES, NO) else "불일치" for j in nat])
    req = crit["injected"]["credit_requires_key"]
    out = {
        "natural_n": len(nat), "natural_agree": agree, "natural_agreement": pct(agree, len(nat)),
        "natural_kappa": kappa, "natural_human_yes": len(human_yes),
        "kappa_used": len(human_yes) >= crit["natural"]["min_positive_for_kappa"],
        "false_alarm": sum(m[j] != NO for j in human_no), "false_alarm_rate": pct(sum(m[j] != NO for j in human_no), len(human_no)),
        "natural_judge_detected": sum(m[j] == YES for j in human_yes),
        "injected_n": len(inj),
        "injected_detected": sum(detected(judged[j], key[j]["keys"], req) for j in inj),
        "injected_detected_any": sum(m[j] == YES for j in inj),
        "rater_injected_detected": sum(h[j] == YES for j in inj),
        "unparsed": sum(not judged[j]["parsed"] for j in key),
        "excluded": [{"번호": j, "item_id": key[j]["item_id"], "human": h[j], "judge": judged[j]["verdict"]}
                     for j in key if key[j]["item_id"] in excluded],
    }
    out["injected_recall"] = pct(out["injected_detected"], len(inj))
    out["injected_phrase_is_whole_summary"] = sum(judged[j]["verdict"] == YES and judged[j]["phrase"] == key[j]["summary"]
                                                  for j in inj)   # 기록 — 요약 전체를 지목해 핵심어 일치가 된 경우
    out["near_boundary"] = (out["natural_agreement"] is not None
                            and abs(out["natural_agreement"] - crit["natural"]["agreement_min"]) <= 3.0)   # 계획 §7
    return out


def decide(mx: dict, crit: dict) -> dict:
    """계획 §6 — 기준값과 같으면 충족."""
    n, i = crit["natural"], crit["injected"]
    checks = [("자연 응답 일치율(%)", mx["natural_agreement"], f"≥ {n['agreement_min']}", mx["natural_agreement"] >= n["agreement_min"]),
              ("주입 날조 검출률(%)", mx["injected_recall"], f"≥ {i['recall_min']}", mx["injected_recall"] >= i["recall_min"]),
              ("거짓 경보율(%)", mx["false_alarm_rate"], f"≤ {n['false_alarm_max']}",
               mx["false_alarm_rate"] is not None and mx["false_alarm_rate"] <= n["false_alarm_max"])]
    if mx["kappa_used"]:
        checks.append(("자연 응답 κ", mx["natural_kappa"], f"≥ {n['kappa_min']}",
                       mx["natural_kappa"] is not None and mx["natural_kappa"] >= n["kappa_min"]))
    checks = [{"name": a, "value": b, "limit": c, "passed": d} for a, b, c, d in checks]
    return {"adopted": all(c["passed"] for c in checks), "checks": checks}


def detector_reference(measure: dict, key: dict) -> dict:
    """기록 — v1.2 규칙 기반 탐지기의 같은 세트 판정."""
    flags = {r["번호"]: detect(f"[요약]: {r['요약']}", source_of(r))["hallucinated"] for r in measure["items"]}
    inj = [j for j, e in key.items() if e["injected"]]
    return {"flags": flags, "injected_detected": sum(flags[j] for j in inj),
            "by_kind": dict(Counter(key[j]["kind"] for j in inj if flags[j]))}


def render(meta: dict, mx: dict, verdict: dict, key: dict, human: dict, judged: dict, det: dict) -> str:
    L = ["# v1.5 LLM-as-judge 판정 — 요약 날조", "",
         f"> 계획 [`judge_v15_plan.md`](judge_v15_plan.md) · 기준 `judge_v15.toml` · `src/triage_eval/pipeline/judge.py` 생성 · "
         f"채점관 `{meta['judge_model']}` (digest {meta['model_digest']}, 프롬프트 {meta['prompt_version']})", "",
         f"## 판정: **{'채택' if verdict['adopted'] else '미채택'}**", "",
         "| 기준 | 값 | 기준 | 결과 |", "| :--- | ---: | :--- | :---: |"]
    L += [f"| {c['name']} | {c['value']} | {c['limit']} | {'✅' if c['passed'] else '❌'} |" for c in verdict["checks"]]
    if not mx["kappa_used"]:
        L += ["", f"> 검토자의 자연 응답 날조 있음이 {mx['natural_human_yes']}건이라 κ({mx['natural_kappa']})는 기록만 했다(계획 §6)."]
    L += ["", "## 기록 (판정에 쓰지 않음)", "",
          f"- 자연 응답 지표에서 뺀 문항(ISSUE-012): " + (", ".join(f"{x['번호']}(검토자 {x['human']}, 채점관 {x['judge']})"
                                                             for x in mx["excluded"]) or "없음"),
          f"- 자연 응답 {mx['natural_n']}건: 일치 {mx['natural_agree']}, 검토자 있음 {mx['natural_human_yes']}건 중 채점관 있음 "
          f"{mx['natural_judge_detected']}건, 거짓 경보 {mx['false_alarm']}건",
          f"- 주입 {mx['injected_n']}건: 채점관 검출(문구 일치) {mx['injected_detected']}, 문구 무관 {mx['injected_detected_any']}, "
          f"검토자 검출 {mx['rater_injected_detected']}, v1.2 탐지기 {det['injected_detected']}",
          f"- 채점관 출력을 읽지 못한 건: {mx['unparsed']}, 요약 전체를 지목한 주입 검출: {mx['injected_phrase_is_whole_summary']}",
          *([f"- 자연 응답 일치율이 기준 ±3%p 안 — 외부 QA 1명이 같은 세트를 채점할 후보(계획 §7)"] if mx["near_boundary"] else []),
          "- 정의는 검토자 시트와 채점관 프롬프트에 같은 문구로 주었다. 채점관 프롬프트에는 판정 방법(진술 단위 확인)이, "
          "검토자 시트에는 가상 예시 2개가 더 있다", "",
          "## 주입 날조별 결과", "", "| 번호 | 종류 | 채점관 | 지목 문구 | 검토자 | v1.2 탐지기 |", "| :--- | :--- | :--- | :--- | :--- | :---: |"]
    for j, e in sorted(key.items()):
        if e["injected"]:
            L.append(f"| {j} | {e['kind']} | {judged[j]['verdict'] or '읽지 못함'} | {judged[j]['phrase'] or '—'} | "
                     f"{human[j]['verdict']} | {'있음' if det['flags'][j] else '—'} |")
    L += ["", "## 자연 응답에서 판정이 갈린 건", "", "| 번호 | 요약 | 검토자 | 채점관 | 채점관 지목 문구 | 검토자 문구 |",
          "| :--- | :--- | :--- | :--- | :--- | :--- |"]
    for j, e in sorted(key.items()):
        if not e["injected"] and human[j]["verdict"] != judged[j]["verdict"]:
            L.append(f"| {j} | {e['summary']} | {human[j]['verdict']} | {judged[j]['verdict'] or '읽지 못함'} | "
                     f"{judged[j]['phrase'] or '—'} | {human[j]['phrase'] or '—'} |")
    return "\n".join(L) + "\n"


def git_lines(*args: str) -> list[str]:
    out = subprocess.run(["git", *args], cwd=ROOT, capture_output=True, check=True)
    return [x for x in out.stdout.decode("utf-8").splitlines() if x.strip()]


def measure_guard(client, crit: dict, sheet: Path) -> str | None:
    """설계 의도 8 — 측정해도 되면 실행 전 모델 digest, 아니면 GateError."""
    problems = []
    if RESULT_JSON.exists():
        problems.append(f"측정 결과가 이미 있습니다({rel(RESULT_JSON)}). 측정은 한 번만 한다")
    dirty = git_lines("status", "--porcelain", "--", *GUARDED)
    if dirty:
        problems.append("커밋하지 않은 변경이 있습니다: " + ", ".join(dirty))
    sheet_path = sheet.resolve()
    if not sheet_path.is_relative_to(ROOT) or not git_lines("ls-files", "--", sheet_path.relative_to(ROOT).as_posix())             or git_lines("status", "--porcelain", "--", sheet_path.relative_to(ROOT).as_posix()):
        problems.append(f"검토자 시트 {sheet.name} 가 수정 없이 저장소에 커밋되어 있지 않습니다(계획 §5-3)")
    dev = json.loads(DEV_RESULT.read_text(encoding="utf-8"))
    digest = get_model_digests(client, [crit["meta"]["judge_model"]]).get(crit["meta"]["judge_model"])
    if dev.get("prompt_sha256") != prompt_fingerprint():
        problems.append("프롬프트 지문이 개발용 실행과 다릅니다 — 확정한 프롬프트가 아닙니다")
    if dev.get("model_digest") != digest:
        problems.append(f"모델 digest가 개발용 실행과 다릅니다 ({dev.get('model_digest')} → {digest})")
    if problems:
        raise GateError("; ".join(problems))
    return digest


def record_invalid(meta: dict, reasons: list[str]) -> None:
    """설계 의도 9 — 무효 실행을 남긴다."""
    runs = json.loads(INVALID_RUNS.read_text(encoding="utf-8")) if INVALID_RUNS.exists() else []
    runs.append({**meta, "reasons": reasons})
    INVALID_RUNS.write_text(json.dumps(runs, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def run_measure(client, crit: dict, sheet: Path, write: bool = True, guard: bool = True) -> int:
    if not write:
        raise GateError("측정 모드는 결과를 반드시 기록한다(--no-write 불가, 설계 의도 8)")
    measure = json.loads((ROOT / crit["meta"]["measure_set"]).read_text(encoding="utf-8"))
    key = json.loads((ROOT / crit["meta"]["answer_key"]).read_text(encoding="utf-8"))["map"]
    human = read_rater_sheet(sheet, list(key), {r["번호"]: r["요약"] for r in measure["items"]})   # 설계 의도 6 — 채점관 실행 전에 확인
    model = crit["meta"]["judge_model"]
    digest = measure_guard(client, crit, sheet) if guard else get_model_digests(client, [model]).get(model)
    results, reloads = run_judge(client, model, measure["items"])
    meta = {**run_meta(client, model, digest), "sheet": sheet.name}
    if reloads:
        record_invalid(meta, reloads)
        raise GateError("측정 환경 오염 — 실행 도중 모델 재로드·통신 오류: " + ", ".join(reloads)
                        + f". {rel(INVALID_RUNS)} 에 기록했다. 계획 §4대로 환경을 정리하고 처음부터 다시 측정한다")
    judged = {r["번호"]: r for r in results}
    mx = metrics(key, human, judged, crit)
    verdict = decide(mx, crit)
    det = detector_reference(measure, key)
    for c in verdict["checks"]:
        print(f"  {'✅' if c['passed'] else '❌'} {c['name']}: {c['value']} ({c['limit']})")
    print(f"판정: {'채택' if verdict['adopted'] else '미채택'}")
    if write:
        RESULT_JSON.write_text(json.dumps({**meta, "metrics": mx, "verdict": verdict, "human": human, "judge": judged,
                                           "detector": det}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        REPORT.write_text(render(meta, mx, verdict, key, human, judged, det), encoding="utf-8")
        print(f"저장: {rel(RESULT_JSON)}, {rel(REPORT)}")
    return 0 if verdict["adopted"] else 1


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="v1.5 LLM-as-judge 요약 날조 판정")
    mode = p.add_mutually_exclusive_group(required=True)
    mode.add_argument("--dev", action="store_true", help="개발용 자료로 채점관 실행 (계획 §5)")
    mode.add_argument("--measure", action="store_true", help="측정 세트로 1회 측정 (계획 §5-4)")
    p.add_argument("--sheet", type=Path, help="검토자가 채점한 시트 (--measure)")
    p.add_argument("--no-write", action="store_true")
    args = p.parse_args(argv)
    import ollama

    try:
        crit = load_criteria()
        if args.dev:
            return run_dev(ollama.Client(), crit, write=not args.no_write)
        if not args.sheet:
            raise GateError("--measure 에는 --sheet 로 검토자가 채점한 시트를 지정해야 합니다")
        return run_measure(ollama.Client(), crit, args.sheet, write=not args.no_write)
    except GateError as e:
        print(f"판정 불가: {e}")
        return 2


if __name__ == "__main__":
    sys.exit(main())
