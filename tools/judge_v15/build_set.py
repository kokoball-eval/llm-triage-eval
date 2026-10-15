"""v1.5 LLM-as-judge — 요약 날조 측정 세트·개발용 자료·정답 대응표·검토자 채점 시트를 만든다.

계획: report/judge_v15_plan.md (§3 측정 세트, §5 개발용 자료) · 기준: judge_v15.toml

사용법
------
    uv run python tools/judge_v15/build_set.py              # data/judge_v15/ 3개 파일(없을 때만) + dist/judge_v15/ 채점 시트
    uv run python tools/judge_v15/build_set.py --out <폴더>  # 채점 시트를 쓸 폴더 지정

출력
----
    data/judge_v15/measure_set_v15.json   측정 세트 63건 (J01~J63, 리포트당 1건. 리포트 원문과 요약만, 출처·주입 여부 없음)
    data/judge_v15/answer_key_v15.json    정답 대응표 (번호 → 문항·실행 기록·주입 여부·주입 문장)
    data/judge_v15/dev_set_v15.json       채점관 프롬프트 개발용 자료 (seed 11 응답, 측정 세트와 같은 요약은 뺌)
    <out>/요약_날조_채점시트.xlsx          검토자 채점 시트 (기본 dist/judge_v15/, 저장소에 넣지 않음)
    검토자는 채점을 마칠 때까지 answer_key·dev_set 과 이 파일을 열지 않는다(주입 여부·주입 문장이 들어 있다).

[설계 의도]
1. 모델을 새로 실행하지 않는다. 저장된 실행 기록의 [요약]만 쓴다. 측정 세트는 seed 1 실행, 개발용 자료는 seed 11 실행에서 만든다.
   같은 실행 기록에서 프롬프트를 다듬고 측정하면 프롬프트가 측정 세트에 맞춰진다(계획 §5).
2. 측정 세트는 리포트 63개를 한 번씩만 넣는다(주입 20 + 자연 43). 처음에는 100건(자연 80 + 주입 20)이었으나, 주입 문항만
   한 번 등장하고 두 번 나오는 리포트는 모두 자연 응답이라 등장 횟수가 주입의 단서가 되었다. 같은 리포트의 요약 두 개를
   비교해 주입 문장을 찾을 여지도 함께 없앤다.
3. 자연 응답은 qwen2.5:7b의 세 구성(M2 채택 구성·M0·M1) seed 1 실행에서 리포트마다 고정 seed로 하나를 고른다. 채점관(gemma4:12b)이
   쓴 요약은 넣지 않는다. 자기가 쓴 요약을 자기가 채점하는 편향을 막기 위해서다.
4. 주입 대상 문항과 종류는 고정 seed로 정했고, 주입 문장은 각 요약의 내용·말투에 맞게 원문에 없는 사실을 요약 문장 안에
   녹여 썼다(INJECTIONS). 처음에는 짧은 구를 쉼표 뒤에 붙였으나 주입 20건 모두에 쉼표가 생겨(자연 응답은 80건 중 1건)
   겉모양으로 드러났다. 숫자도 단서가 되므로 기기 주입은 숫자 없는 표현을 쓴다. 원래 요약부터 원문과 어긋나는 문항과
   주입 사실이 원문의 함의로 읽힐 수 있는 문항(처음 고른 대상 중 3건)은 대상에서 뺐다.
5. 만들 때마다 확인한다: 원래 요약이 실행 기록과 같은가, 핵심어가 주입 요약에는 있고 그 문항의 리포트 원문(게시판·게시 일시·
   제목·본문·공지·기존 이슈)과 원래 요약에는 없는가. 핵심어가 원문에 있으면 날조가 아니다.
6. 개발용 자료에서는 측정 세트와 글자까지 같은 (리포트, 요약)을 뺀다. 개발용 주입 대상은 측정용 주입 문항과 겹치지 않는다.
   리포트 63개를 둘 다 쓰는 것은 피할 수 없으므로 계획 §8에 한계로 적는다.
7. 측정 세트·대응표·개발용 자료는 봉인한다. 파일이 이미 있고 내용이 다르면 쓰지 않고 종료 코드 2로 끝낸다.
8. 정답 라벨은 쓰지 않는다. 날조 판정은 원문과 요약만으로 하므로 라벨이 필요 없고, 라벨 쪽 정보가 섞이지 않게 한다.
"""

import argparse
import json
import random
from pathlib import Path

from triage_eval.common.paths import ROOT
from triage_eval.pipeline.contract import parse_v2
from triage_eval.pipeline.prompt import TRACK_LABEL

SEED = 1515
DATASET = ROOT / "data" / "eval_v14" / "aether_raid_v14.json"   # 실행 기록을 만든 평가셋 (입력은 v1.5와 같다)
OUT_DIR = ROOT / "data" / "judge_v15"
FILES = ("measure_set_v15.json", "answer_key_v15.json", "dev_set_v15.json")
DEFAULT_SHEET_DIR = ROOT / "dist" / "judge_v15"
H = "data/results/v14/history"
ADOPTED = ("data/results/v14/baseline/v14_dev_m2_seed1_baseline.json", "qwen2.5:7b")    # 채택 구성 seed 1
NATURAL_SOURCES = [ADOPTED,                                                              # 설계 의도 3
                   (f"{H}/v14_dev_m0_20261004_204024.json", "qwen2.5:7b"),               # M0 seed 1
                   (f"{H}/v14_dev_m1_20261004_204436.json", "qwen2.5:7b")]               # M1 seed 1
DEV = ("data/results/v14/baseline/v14_dev_m2_seed11_baseline.json", "qwen2.5:7b")       # 채택 구성 seed 11
JUDGE_MODEL = "gemma4:12b"

# 설계 의도 4·5 — {문항 ID: (종류, 원래 요약, 주입한 요약, 핵심어)}. 측정용 20건(기기4·수치4·시점3·추가 현상5·원인 단정4)
INJECTIONS = {
    "A08": ("시점", "4장 엔딩 컷신에서 세라의 얼굴 텍스처가 이상하게 표시됨",
            "어제 밤 패치 이후 4장 엔딩 컷신에서 세라의 얼굴 텍스처가 이상하게 표시됨", ["어제 밤"]),
    "A11": ("기기", "로그인 문제, 상점 가격 오류, 친구 목록 문제 발생",
            "듀얼 모니터 PC에서 로그인 문제, 상점 가격 오류, 친구 목록 문제 발생", ["듀얼 모니터"]),
    "A12": ("수치", "필드에서 몹 모이면 프레임 저하, 스킬 효과음 지연, 이벤트 창에서 영어 문구 표시",
            "필드에서 몹 30마리 이상 모이면 프레임 저하, 스킬 효과음 지연, 이벤트 창에서 영어 문구 표시", ["30마리"]),
    "A14": ("추가 현상", "PC에서 추석 한정 다이아 패키지 결제 후 다이아가 미지급",
            "PC에서 추석 한정 다이아 패키지 결제 후 다이아가 미지급되고 결제 내역도 사라짐", ["결제 내역"]),
    "A18": ("추가 현상", "10연차에서 SSR 두 장 뽑음", "10연차에서 SSR 두 장 뽑은 뒤 뽑기 화면이 멈춤", ["화면이 멈춤"]),
    "B02": ("기기", "메인 퀘스트 4-7에서 NPC 리안이 스폰되지 않아 퀘스트 진행 불가",
            "아이폰에서 메인 퀘스트 4-7 NPC 리안이 스폰되지 않아 퀘스트 진행 불가", ["아이폰"]),
    "B07": ("수치", "아이템 툴팁 설명 텍스트가 툴팁 상자 바깥으로 넘침",
            "아이템 툴팁 설명 텍스트가 툴팁 상자 바깥으로 약 20px 넘침", ["20px"]),
    "B09": ("원인 단정", "캐릭터가 특정 벽면에 붙을 때 팔이 벽 안으로 들어갑니다",
            "충돌 판정 데이터 누락으로 캐릭터가 특정 벽면에 붙을 때 팔이 벽 안으로 들어갑니다", ["충돌 판정"]),
    "B10": ("시점", "메인 스토리 4-3 컷신 대사 오타", "주말 이벤트 시작 후 메인 스토리 4-3 컷신 대사 오타", ["주말"]),
    "B11": ("원인 단정", "분수대 그림자 미출력", "그림자 셰이더 누락으로 분수대 그림자 미출력", ["셰이더"]),
    "B15": ("수치", "PC 계정 연동 후 모바일 데이터 사라짐", "PC 계정 연동 후 유저 200여 명의 모바일 데이터 사라짐", ["200"]),
    "B16": ("원인 단정", "유저가 매크로 사용자 신고를 하였으나 제재가 이루어지지 않았다",
            "유저가 매크로 사용자 신고를 하였으나 신고 접수 서버 오류로 제재가 이루어지지 않았다", ["서버 오류"]),
    "B18": ("추가 현상", "펼친 화면에서 전투 UI 우측 스킬 버튼이 잘려 터치 불가",
            "펼친 화면에서 전투 UI 우측 스킬 버튼과 미니맵이 잘려 터치 불가", ["미니맵"]),
    "S03": ("추가 현상", "보스전 맵 구석의 뾰족한 바위 위에 점프하면 보스가 공격을 멈춥니다",
            "보스전 맵 구석의 뾰족한 바위 위에 점프하면 보스가 공격을 멈추고 체력도 회복됩니다", ["회복"]),
    "S04": ("원인 단정", "랭크전에서 패배 기록이 남지 않는 현상 발생", "서버 부하로 랭크전에서 패배 기록이 남지 않는 현상 발생", ["부하"]),
    "S09": ("수치", "갤럭시 탭 S4에서 PC 계정으로 로그인 시 앱이 종료됨",
            "갤럭시 탭 S4에서 PC 계정으로 로그인 시 앱이 3초 만에 종료됨", ["3초"]),
    "S10": ("기기", "사용자가 업데이트 후 스킬 데미지가 줄어든 것에 대해 버그로 판단",
            "사용자가 업데이트 후 PC에서만 스킬 데미지가 줄어든 것에 대해 버그로 판단", ["PC에서만"]),
    "S16": ("추가 현상", "길드전 중 길드원 전원이 동시에 튕김 발생", "길드전 중 길드원 전원이 동시에 튕기고 길드전 보상도 지급되지 않음", ["지급되지"]),
    "S18": ("시점", "같은 템의 경우 다른 플레이어에게 훨씬 높은 데미지를 입힌다",
            "지난주부터 같은 템의 경우 다른 플레이어에게 훨씬 높은 데미지를 입힌다", ["지난주"]),
    "S20": ("기기", "파티 4명이 물약 다 써도 보스를 잡지 못함", "폴더블 기기에서 파티 4명이 물약 다 써도 보스를 잡지 못함", ["폴더블"]),
}
# 개발용 10건 — 측정용 주입 문항과 겹치지 않고, seed 11 요약이 측정 세트의 요약과 다른 문항 (설계 의도 6)
DEV_INJECTIONS = {
    "A02": ("원인 단정", "보스방에 들어가면 튕김 발생", "패치 다운로드 오류로 보스방에 들어가면 튕김 발생", ["다운로드"]),
    "A04": ("추가 현상", "던전 진입 시 끊김 발생", "던전 진입 시 끊김 발생 후 인벤토리 아이템 일부 소실", ["인벤토리"]),
    "A09": ("수치", "신규 보스 난이도가 높아 파티 플레이가 어려움", "신규 보스 난이도가 높아 파티 클리어율이 5%에 불과", ["5%"]),
    "B14": ("기기", "폴드 펼친 화면에서 스킬 버튼이 화면 밖으로 잘림", "태블릿과 폴드 펼친 화면에서 스킬 버튼이 화면 밖으로 잘림", ["태블릿"]),
    "B17": ("원인 단정", "전투 중 전화 수신 후 복귀 시 사운드 끊김과 화면 정지 현상 발생",
            "메모리 누수로 전투 중 전화 수신 후 복귀 시 사운드 끊김과 화면 정지 현상 발생", ["메모리"]),
    "S01": ("추가 현상", "경매장 수령액이나 우편함 보상을 여러 번 중복으로 타먹는 애들에 대한 제보",
            "경매장 수령액이나 우편함 보상을 여러 번 중복으로 타먹고 계정을 삭제하는 애들에 대한 제보", ["계정을 삭제"]),
    "S13": ("시점", "55레벨로 밀려난다", "오늘 오전 7시경 55레벨로 밀려난다", ["7시"]),
    "S14": ("수치", "세 번째 레이드 보상이 안 들어온다", "세 번째 레이드 보상 중 강화석 15개가 안 들어온다", ["15개"]),
    "S15": ("추가 현상", "신규 레이드 보스 싸우면 앱 꺼짐, 보상 없음", "신규 레이드 보스 싸우면 앱 꺼지고 음성 채팅도 끊김, 보상 없음", ["음성"]),
    "S17": ("추가 현상", "추석 픽업 100회 돌렸으나 SSR 픽업 SSR이 나온 횟수 미달",
            "길드원 전원이 추석 픽업 100회 돌렸으나 SSR 픽업 SSR이 나온 횟수 미달", ["길드원"]),
}


# ── 읽기 ─────────────────────────────────────────────────────

def load_items() -> dict[str, dict]:
    return {i["id"]: i for i in json.loads(DATASET.read_text(encoding="utf-8"))["items"] if i["split"] == "dev"}


def summaries(log_path: str, model: str, run_index: int = 1) -> dict[str, str]:
    """{문항 ID: 요약}. 요약을 읽지 못한 응답은 뺀다."""
    log = json.loads((ROOT / log_path).read_text(encoding="utf-8"))
    out = {}
    for r in log["results"]:
        if r["model"] == model and r["run_index"] == run_index and r.get("success"):
            s = (parse_v2(r["response_text"]).get("요약") or "").strip()
            if s:
                out[r["item_id"]] = s
    return out


def source_text(item: dict) -> str:
    """설계 의도 5 — 날조 판정의 근거가 되는 리포트 원문 전체(검토자 시트·채점관 입력에 보이는 모든 칸)."""
    inp, ctx = item["input"], item["input"].get("context", {})
    return "\n".join([inp.get("board", ""), inp.get("reporter", ""), inp.get("posted_at", ""), inp["title"], inp["body"],
                      ctx.get("update_note", ""), *ctx.get("known_issues", [])])


def check_injections(injections: dict, items: dict, base: dict[str, str]) -> None:
    """설계 의도 5 — 원래 요약이 실행 기록과 같고, 핵심어가 주입 요약에는 있고 원문·원래 요약에는 없어야 한다."""
    problems = []
    for i, (_, original, injected, keys) in injections.items():
        if base.get(i) != original:
            problems.append(f"{i} 원래 요약이 실행 기록과 다릅니다")
        for k in keys:
            if k not in injected:
                problems.append(f"{i} 핵심어 {k!r} 가 주입 요약에 없습니다")
            if k in source_text(items[i]) or k in original:
                problems.append(f"{i} 핵심어 {k!r} 가 원문이나 원래 요약에 있어 날조가 아닙니다")
    if problems:
        raise ValueError("; ".join(problems))


# ── 세트 만들기 ──────────────────────────────────────────────

def report_fields(item: dict) -> dict:
    inp, ctx = item["input"], item["input"].get("context", {})
    a = item["track"] == "A"
    return {"입력 종류": TRACK_LABEL[item["track"]], "게시판 / 작성자": inp["board"] if a else inp["reporter"],
            "게시 일시": inp["posted_at"] if a else "", "제목": inp["title"], "본문": inp["body"],
            "참고: 업데이트 공지": ctx.get("update_note", ""),
            "참고: 기존 등록 이슈": "\n".join(f"- {k}" for k in ctx.get("known_issues", []))}


def build(seed: int = SEED) -> tuple[dict, dict, dict]:
    """(측정 세트, 정답 대응표, 개발용 자료)."""
    rng = random.Random(seed)
    items = load_items()
    sources = {src: summaries(*src) for src in NATURAL_SOURCES}
    adopted = sources[ADOPTED]
    if sorted(adopted) != sorted(items):
        raise ValueError(f"채택 구성 seed 1 1회차 요약이 {len(adopted)}건입니다 (개발용 {len(items)}건이어야 함)")
    check_injections(INJECTIONS, items, adopted)
    entries = []
    for i in sorted(items):
        if i in INJECTIONS:
            kind, original, injected, keys = INJECTIONS[i]
            entries.append({"item_id": i, "log": ADOPTED[0], "model": ADOPTED[1], "injected": True, "kind": kind,
                            "keys": keys, "original_summary": original, "summary": injected})
        else:   # 설계 의도 3 — 요약이 있는 구성 중에서 고정 seed로 하나
            src = rng.choice([s for s in NATURAL_SOURCES if i in sources[s]])
            entries.append({"item_id": i, "log": src[0], "model": src[1], "injected": False, "summary": sources[src][i]})
    rng.shuffle(entries)
    key, measure = {}, []
    for n, e in enumerate(entries, start=1):
        jid = f"J{n:02d}"
        key[jid] = e
        measure.append({"번호": jid, **report_fields(items[e["item_id"]]), "요약": e["summary"]})

    dev_base = summaries(*DEV)
    check_injections(DEV_INJECTIONS, items, dev_base)
    in_measure = {e["item_id"]: e["summary"] for e in entries}
    dev_ids = [i for i in sorted(dev_base) if dev_base[i] != in_measure.get(i)]           # 설계 의도 6
    bad = sorted(set(DEV_INJECTIONS) - set(dev_ids)) + sorted(set(DEV_INJECTIONS) & set(INJECTIONS))
    if bad:
        raise ValueError(f"개발용 주입 대상 {bad} 가 측정 세트와 겹칩니다")
    dev = [{"item_id": i, "log": DEV[0], "model": DEV[1], "injected": i in DEV_INJECTIONS,
            **({"kind": DEV_INJECTIONS[i][0], "keys": DEV_INJECTIONS[i][3], "original_summary": dev_base[i]}
               if i in DEV_INJECTIONS else {}),
            "summary": DEV_INJECTIONS[i][2] if i in DEV_INJECTIONS else dev_base[i],
            **report_fields(items[i])} for i in dev_ids]

    meta = {"plan": "report/judge_v15_plan.md", "seed": seed, "dataset": DATASET.relative_to(ROOT).as_posix()}
    return ({**meta, "count": len(measure), "items": measure},
            {**meta, "note": "검토자는 채점을 마칠 때까지 이 파일을 열지 않는다 (계획 §3)", "map": key},
            {**meta, "note": "채점관 프롬프트 개발용 (계획 §5). 측정 세트와 다른 실행 기록(seed 11), 측정 세트와 같은 요약은 뺌. "
                             "검토자는 채점을 마칠 때까지 열지 않는다", "items": dev})


def serialize(payload: dict) -> str:
    return json.dumps(payload, ensure_ascii=False, indent=2) + "\n"


def write_sealed(payloads: dict[str, dict]) -> list[str]:
    """설계 의도 7 — 없으면 쓰고, 있으면 내용이 같을 때만 통과. 다른 파일 이름 목록을 돌려준다(쓰지 않음)."""
    changed = [n for n, p in payloads.items() if (OUT_DIR / n).exists()
               and json.loads((OUT_DIR / n).read_text(encoding="utf-8")) != json.loads(serialize(p))]
    if changed:
        return changed
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for n, p in payloads.items():
        if not (OUT_DIR / n).exists():
            (OUT_DIR / n).write_text(serialize(p), encoding="utf-8")
    return []


# ── 채점 시트 ────────────────────────────────────────────────

SHEET_COLS = ["번호", "입력 종류", "게시판 / 작성자", "게시 일시", "제목", "본문", "참고: 업데이트 공지", "참고: 기존 등록 이슈",
              "요약", "판정", "날조 문구", "확신도", "메모"]
VERDICTS = ["날조 있음", "날조 없음"]
CONFIDENCE = ["높음", "보통", "낮음"]
GUIDE = [
    "요약 날조 채점 시트 — 검토자용",
    "",
    "각 행의 [요약]이 왼쪽 리포트 원문(게시판·게시 일시·제목·본문·업데이트 공지·기존 등록 이슈)에 맞는지 판정합니다.",
    "· 요약에 원문에 없거나 원문과 어긋나는 사실 진술이 하나라도 있으면 '날조 있음'",
    "· 원문을 다른 말로 바꿔 쓰거나 줄인 것은 날조가 아님",
    "· '~로 추정', '~일 수 있음'처럼 추정임을 밝힌 표현은 날조가 아님",
    "· '날조 있음'이면 해당 문구를 '날조 문구' 칸에 적어 주세요",
    "",
    "예시(측정 세트와 무관한 가상 리포트)",
    "· 원문 '낚시하면 화면 멈춰요' / 요약 '낚시 중 화면 멈춤' → 날조 없음",
    "· 원문 '낚시하면 화면 멈춰요' / 요약 '낚시 중 화면 멈춤 후 재접속 불가' → 날조 있음 (문구: 재접속 불가)",
    "",
    "모든 행의 '판정'을 드롭다운에서 골라 주세요. 빈칸이 있으면 계산하지 않습니다.",
    "채점을 마칠 때까지 data/judge_v15/ 의 answer_key_v15.json, dev_set_v15.json 은 열지 않습니다.",
]


def write_sheet(measure: dict, out_dir: Path) -> Path:
    from openpyxl import Workbook                      # 개발 의존성 — 시트를 만들 때만 불러온다
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    g = wb.active
    g.title = "시작하기"
    for r, line in enumerate(GUIDE, start=1):
        g.cell(row=r, column=1, value=line).font = Font(name="맑은 고딕", size=14 if r == 1 else 10, bold=r == 1)
    g.column_dimensions["A"].width = 120
    ws = wb.create_sheet("판정")
    head = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    for c, name in enumerate(SHEET_COLS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font, cell.fill = head, PatternFill("solid", fgColor="1F3A5F")
    wrap = Alignment(wrap_text=True, vertical="top")
    for r, row in enumerate(measure["items"], start=2):
        for c, name in enumerate(SHEET_COLS, start=1):
            cell = ws.cell(row=r, column=c, value=row.get(name, ""))
            cell.alignment = wrap
            if name in ("판정", "날조 문구", "확신도", "메모"):
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
            elif name == "요약":
                cell.fill = PatternFill("solid", fgColor="E2EFDA")
        ws.row_dimensions[r].height = min(15 * max(3, row["본문"].count("\n") + 1 + len(row["본문"]) // 55), 409)
    last = len(measure["items"]) + 1
    for name, choices in (("판정", VERDICTS), ("확신도", CONFIDENCE)):
        col = chr(ord("A") + SHEET_COLS.index(name))
        dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        dv.add(f"{col}2:{col}{last}")
        ws.add_data_validation(dv)
    widths = {"번호": 6, "입력 종류": 12, "게시판 / 작성자": 12, "게시 일시": 11, "제목": 24, "본문": 55,
              "참고: 업데이트 공지": 20, "참고: 기존 등록 이슈": 26, "요약": 36, "판정": 11, "날조 문구": 24, "확신도": 8, "메모": 24}
    for c, name in enumerate(SHEET_COLS, start=1):
        ws.column_dimensions[chr(ord("A") + c - 1)].width = widths[name]
    ws.freeze_panes = "B2"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "요약_날조_채점시트.xlsx"
    wb.save(path)
    return path


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="요약 날조 측정 세트·대응표·개발용 자료·채점 시트 생성")
    p.add_argument("--out", type=Path, default=DEFAULT_SHEET_DIR, help="채점 시트를 쓸 폴더")
    args = p.parse_args(argv)
    measure, key, dev = build()
    changed = write_sealed(dict(zip(FILES, (measure, key, dev))))
    if changed:
        print(f"봉인된 파일과 내용이 다릅니다 — 쓰지 않습니다: {', '.join(changed)}")
        return 2
    sheet = write_sheet(measure, args.out)
    print(f"측정 세트 {measure['count']}건 (주입 {sum(v['injected'] for v in key['map'].values())}건), "
          f"개발용 {len(dev['items'])}건 (주입 {sum(d['injected'] for d in dev['items'])}건)")
    print(f"저장: data/judge_v15/ ({', '.join(FILES)}), {sheet}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
